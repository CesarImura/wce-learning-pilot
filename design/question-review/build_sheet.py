import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

steps = json.load(open(sys.argv[1]))
OUT = sys.argv[2]

LABELS = {
    "candle-scrub": "Candle anatomy", "candle-builder": "Build a candle", "tap-engulfing": "Spot patterns",
    "draw-support": "Draw support", "sort-charts": "Trend vs. range", "ma-tuner": "Tune indicators",
    "order-book": "Order book", "predict-fakeout": "Predict & reveal", "place-trade": "Plan a trade",
    "position-sizer": "Position sizing", "expectancy-lab": "Expectancy", "trade-journal": "Trade journal",
    "concept": "Concept", "classify": "Sort into buckets", "order": "Put in order", "directions": "Predict direction",
    "number": "Calculate", "returns-reality": "Returns reality check", "breakeven-curve": "Break-even curve",
    "prop-firm": "Prop-firm payout", "yield-curve": "Yield curve", "dividend-split": "Dividends",
    "fed-cycle": "Central bank decisions", "carry-trade": "Carry trade", "real-yield-gold": "Real yields vs. gold",
    "volume-profile": "Volume profile", "orb-breakout": "Opening range breakout",
}

# Interactive exercises keep their answer and feedback in code (src/components/exercises/*).
INTERACTIVE = {
    "candle-scrub": ("Drag through the hour and watch the open, high, low and close form one candle.",
                     "Not graded: it walks through how one hour becomes a candle.", "", ""),
    "candle-builder": ("Drag open, high, low and close to build the candle for the hour shown.",
                       "The candle matching the hour's open, high, low and close (green because it closed above the open).",
                       "That's exactly the candle for this hour. The body is green because it closed above the open.",
                       "Check your [open / high / low / close]. The dashed outline shows the real candle."),
    "tap-engulfing": ("Tap one candle on a chart that falls into a small red candle followed by a big green one.",
                      "The green candle right after the small red one at the bottom of the downtrend.",
                      "The green body completely swallows the previous red body, right at the end of a downtrend. (\"Why?\" shows an illustrated zoom of the two candles.)",
                      "Next to it: \"So close! The engulfing candle is the big green one that wraps the small red one.\" Elsewhere: \"Look for a small red candle followed by a green candle whose body covers it completely.\""),
    "draw-support": ("Drag a horizontal line onto the support level of the chart.",
                     "Around 100, where price bounced three times.",
                     "Price bounced off ~100 three times. Each touch makes the level more significant.",
                     "Support is the floor price keeps bouncing from. Look for the lows that line up."),
    "sort-charts": ("Label each chart as trending or consolidating (ranging).",
                    "Each chart in its right group (trend = higher highs and higher lows, or the reverse; range = returns to the same levels).",
                    "Trends make higher highs and higher lows (or the reverse). Ranges keep returning to the same levels.",
                    "Chart [n] is misclassified."),
    "ma-tuner": ("Slide the moving-average period until the chart shows 3 or fewer crossover signals.",
                 "Any period that leaves 3 or fewer crossovers.",
                 "With a [n]-period MA you only get [n] signals. The trade-off: longer MAs react later, so you enter later.",
                 "Still [n] crossovers, every one a potential losing trade. Try a longer period to smooth out the noise."),
    "order-book": ("Set the size of a market buy on a BTC order book.",
                   "The largest size whose average fill stays less than 0.03% above the best ask.",
                   "[size] BTC is the most you can buy before the average fill drifts more than 0.03% above the best ask.",
                   "Too small: \"[size] BTC is within the limit, but you can buy more.\" Too big: \"At [size] BTC you're eating into expensive levels: slippage is [x]%.\""),
    "predict-fakeout": ("Price wicks above resistance at 108. Choose: buy, wait, or short.",
                        "Wait: the candle closed back inside the range, so the rule (buy only after a close above resistance) isn't met.",
                        "The candle poked above 108 but closed back inside the range. No close above resistance means no breakout, and price fell back.",
                        "Buy: \"The wick went above 108, but the candle closed back inside. That's a fakeout, and longs got trapped.\" Short: \"It worked this time, but shorting a wick is a guess, not your rule.\""),
    "place-trade": ("Drag the stop loss and take profit on a breakout, then watch the replay.",
                    "Stop below the broken resistance (108), target at least 2× the risk, stop under entry and target above.",
                    "Solid plan. [outcome of the replay] A good plan is judged by its rules, not by one outcome.",
                    "[outcome] The plan broke a rule: [the rules that failed]."),
    "position-sizer": ("Use the calculator to size a trade on a $10,000 account risking 2% with a 40-pip stop.",
                       "0.5 lots.",
                       "0.5 lots × 40 pips × $10 = $200, which is 2% of $10,000.",
                       "Set the stop to 40 pips, then raise risk until the size reaches 0.5 lots. Or: 0.5 × 40 × $10 = $200 = 2%."),
    "expectancy-lab": ("Set win rate and reward-to-risk to find a system that wins under 40% of the time with expectancy of at least +0.2R.",
                       "Any win rate below 40% with an R:R that gives expectancy ≥ +0.2R (e.g. 35% at 1:3).",
                       "[win]% winners at 1:[rr] gives +[e]R per trade. You can be wrong most of the time and still make money.",
                       "\"Bring the win rate below 40%, then find the R:R that still makes the edge positive.\" or \"Expectancy is [e]R. With so few winners, each win needs to pay more.\""),
    "trade-journal": ("Ten trades in R: +2, −1, −1, +3, −1, +2, −1, −1, +2.5, −1. What's the expectancy per trade?",
                      "+0.35R per trade (total +3.5R over 10 trades, 40% win rate).",
                      "Total +3.5R over 10 trades = +0.35R per trade, with only a 40% win rate.",
                      "Add up all the R results (+3.5R), then divide by the number of trades (10)."),
    "returns-reality": ("With $10,000 and Jim Simons-level returns (60% a year), how much do you make in a year?",
                        "$6,000",
                        "Even matching the best quant fund in history turns $10,000 into a $6,000 profit. Nobody lives on that, which is why trading isn't a get-rich-quick business model.",
                        "Shows the right answer and the same explanation. Hint: set the account size from the question, then read the yearly profit of each investor."),
    "breakeven-curve": ("Drag a point onto the break-even curve for a 1:2 strategy.",
                        "1 : 2 at a 33% win rate.",
                        "At 1:2 you only need to win 1 / (1 + 2) = 33% of the time to break even. Anything above the line has a profit factor above 1.",
                        "Shows the right point (1:2 at 33%)."),
    "prop-firm": ("You pass a $200,000 challenge and make 10%. With an 80% split, what's your payout?",
                  "$16,000",
                  "$200,000 × 10% = $20,000 profit, and an 80% split pays you $16,000. Big upside, but most challenge buyers fail, and those fees fund the few payouts.",
                  "Shows the right answer and the same explanation."),
    "yield-curve": ("Drag the yields of each maturity (3M to 30Y) into a normal yield curve.",
                    "Each maturity pays at least as much as the shorter one, and 30Y pays at least 1% more than 3M.",
                    "That's a normal yield curve: the longer you lend, the more you get paid (the term premium). When short yields rise above long ones, the curve inverts, typically when rate hikes are expected.",
                    "Each maturity should pay at least as much as the shorter one before it, and 30Y should pay clearly more than 3M (at least 1% more)."),
    "dividend-split": ("The company reinvests $100k of its $500k earnings. How much does founder C receive?",
                       "$200,000",
                       "$500k − $100k reinvested = $400k paid out. C owns 50%, so C receives $200k. Shares give you ownership and a cut of the profits.",
                       "Shows the right answer and the same explanation."),
    "fed-cycle": ("At each point of the cycle, choose cut, hold or hike rates.",
                  "Cut when unemployment is the problem, hike when inflation is the problem.",
                  "Cut when unemployment is the problem (cheap credit → more borrowing → more spending). Hike when inflation is the problem. And notice markets move 6–9 months ahead of the economy.",
                  "Central banks cut rates to fight unemployment and hike rates to fight inflation. Low rates make borrowing (and new money) cheaper."),
    "carry-trade": ("The Fed is at 4.5% and the Bank of Japan at 0.5%. What does the $100,000 carry trade earn per year?",
                    "$4,000",
                    "4.5% − 0.5% = 4% on $100,000 = $4,000 a year for borrowing cheap and lending dear. That flow into dollars is why USD/JPY trends when rate gaps are wide.",
                    "Shows the right answer and the same explanation."),
    "real-yield-gold": ("A bond yields 5% and inflation is 2%. What's the real yield?",
                        "3%",
                        "5% − 2% = a 3% real yield. That's a solid, safe inflation hedge, so money prefers bonds and gold tends to suffer. When real yields fall, gold shines.",
                        "Shows the right answer and the same explanation."),
    "volume-profile": ("Tap the point of control on a volume profile.",
                       "The single longest volume bar (the POC).",
                       "The point of control is the price where the most money changed hands. The shaded value area holds 70% of the volume, and the thin zone between the two balances is a low-volume node.",
                       "The POC is the single longest bar: where most trading happened. Price spent the most time there, in balance."),
    "orb-breakout": ("Tap the candle that triggers the opening range breakout.",
                     "The first 5-minute candle that closes above the opening-range high, on a volume spike.",
                     "First 5-minute close above the opening range, on a volume spike. Entry [x], stop below the range at [y], target at 1:2 = [z].",
                     "Inside the range: \"Those candles are the opening range itself.\" Wick only: \"That candle wicked above the range but closed back inside.\" Other: \"Look for the first candle whose close is above the opening-range high, with volume picking up.\""),
}


def fmt(step):
    """(question / content, correct answer, explanation, wrong feedback) for one step."""
    k, p = step["kind"], step["props"] or {}
    if k == "concept":
        parts = []
        if p.get("quote"):
            parts.append(f"Quote: “{p['quote']}”")
        parts += p.get("body", [])
        for pt in p.get("points", []) or []:
            parts.append(f"• {pt['title']}: {pt['text']}")
        return "\n\n".join(parts), "Not graded (concept screen)", "", ""
    if k == "classify":
        buckets = {b["id"]: b["label"] for b in p["buckets"]}
        q = "Buckets: " + "; ".join(f"{b['label']} ({b['hint']})" if b.get("hint") else b["label"] for b in p["buckets"])
        q += "\nItems: " + "; ".join(i["label"] for i in p["items"])
        a = "\n".join(f"{i['label']} → {buckets[i['bucket']]}" for i in p["items"])
        return q, a, p.get("explain", ""), "Shows which items were misplaced, plus the explanation."
    if k == "order":
        q = f"Put in order from {p.get('topLabel', 'top')} to {p.get('bottomLabel', 'bottom')}: " + "; ".join(p["items"])
        a = " → ".join(p["items"])
        return q, a, p.get("explain", ""), "Shows how many are out of place, plus the explanation."
    if k == "directions":
        q = (p.get("context", "") + "\n" if p.get("context") else "") + "\n".join(r["label"] for r in p["rows"])
        up, down = p.get("upLabel", "Up"), p.get("downLabel", "Down")
        a = "\n".join(f"{r['label']}: {up if r['answer'] == 'up' else down} ({r['why']})" for r in p["rows"])
        return q, a, p.get("explain", ""), "\"[n] of [total] wrong.\" plus the explanation."
    if k == "number":
        q = p.get("question", "")
        if p.get("given"):
            q += "\nGiven: " + "; ".join(f"{g['label']}: {g['value']}" if isinstance(g, dict) else str(g) for g in p["given"])
        if p.get("sequence"):
            q += "\nSequence: " + " ".join(s["label"] for s in p["sequence"])
        ans = f"{p.get('unitBefore', '')}{p['answer']}{p.get('unit', '')}"
        if p.get("tolerance"):
            ans += f" (accepts ±{p['tolerance']})"
        return q, ans, p.get("correct", ""), ("Hint: " + p["hint"]) if p.get("hint") else ""
    q, a, right, wrong = INTERACTIVE[k]
    return q, a, right, wrong


# One row per question; unit-test repeats are folded into "Also in".
rows, index = [], {}
counters = {}
for s in steps:
    key = (s["kind"], s["prompt"], json.dumps(s["props"], sort_keys=True))
    if key in index:
        index[key]["also"].append(f"{s['course']} › {s['lesson']}")
        continue
    q, a, right, wrong = fmt(s)
    help_ = s.get("help") or {}
    links = "\n".join(f"{l['label']} ({l['source']}): {l['url']}" for l in help_.get("links", []) or [])
    row = dict(s, q=q, a=a, right=right, wrong=wrong, hint=help_.get("hint", ""), links=links, also=[])
    index[key] = row
    rows.append(row)

wb = Workbook()
ws = wb.active
ws.title = "Review"
cols = [
    ("ID", 9), ("Course", 16), ("Unit", 24), ("Lesson", 24), ("Step", 6), ("Type", 18), ("Prompt (title on screen)", 38),
    ("Question / content", 60), ("Correct answer", 40), ("Explanation (after a right answer)", 50),
    ("Wrong-answer feedback / hint", 40), ("\"Need a hint?\"", 34), ("Further reading", 34), ("Also used in", 30),
    ("Status", 13), ("Reviewer comment", 40), ("Suggested rewrite", 40),
]
head_fill = PatternFill("solid", fgColor="1F2933")
review_fill = PatternFill("solid", fgColor="FFF8E1")
thin = Side(style="thin", color="D9D9D9")
for c, (name, width) in enumerate(cols, 1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = width
ws.row_dimensions[1].height = 32

abbr = {"beginner-series": "BS", "trading-fundamentals": "TF", "order-flow": "OF", "trader-dna": "DNA"}
for r, row in enumerate(rows, 2):
    pre = abbr.get(row["courseId"], row["courseId"][:3].upper())
    counters[pre] = counters.get(pre, 0) + 1
    typ = LABELS.get(row["kind"], row["kind"])
    if row["mode"] == "learn":
        typ += " (learn)"
    vals = [
        f"{pre}-{counters[pre]:03d}", row["course"], row["unit"], row["lesson"] + (" (unit test)" if row["isTest"] and "unit test" not in row["lesson"].lower() else ""),
        row["n"], typ, row["prompt"], row["q"], row["a"], row["right"], row["wrong"], row["hint"], row["links"],
        "\n".join(row["also"]), "To review", "", "",
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
        if c >= 15:
            cell.fill = review_fill

last = len(rows) + 1
dv = DataValidation(type="list", formula1='"To review,OK,Change,Remove"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"O2:O{last}")
for status, color in (("OK", "D7F5E3"), ("Change", "FFE7C2"), ("Remove", "FAD4D4")):
    ws.conditional_formatting.add(f"O2:O{last}", FormulaRule(formula=[f'$O2="{status}"'], fill=PatternFill("solid", fgColor=color)))
ws.freeze_panes = "H2"
ws.auto_filter.ref = f"A1:Q{last}"

# Summary
sm = wb.create_sheet("Summary")
sm.append(["Course", "Questions", "Concept screens", "Graded", "OK", "Change", "Remove", "To review"])
for c in sm[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
courses = []
for row in rows:
    if row["course"] not in courses:
        courses.append(row["course"])
for i, course in enumerate(courses, 2):
    n = sum(1 for x in rows if x["course"] == course)
    concepts = sum(1 for x in rows if x["course"] == course and x["kind"] == "concept")
    sm.append([course, n, concepts, n - concepts,
               f'=COUNTIFS(Review!B:B,A{i},Review!O:O,"OK")', f'=COUNTIFS(Review!B:B,A{i},Review!O:O,"Change")',
               f'=COUNTIFS(Review!B:B,A{i},Review!O:O,"Remove")', f'=COUNTIFS(Review!B:B,A{i},Review!O:O,"To review")'])
t = len(courses) + 2
sm.append(["Total"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{t-1})" for c in range(2, 9)])
for c in sm[t]:
    c.font = Font(bold=True)
for c, w in zip("ABCDEFGH", (26, 11, 16, 9, 8, 9, 9, 11)):
    sm.column_dimensions[c].width = w

# How to review
hw = wb.create_sheet("How to review", 0)
lines = [
    ("WCE Training · question review", True),
    ("", False),
    ("Every question in the Training prototype, one row each. A question used in more than one place appears once; \"Also used in\" lists the other lessons. Trading Fundamentals, Order Flow and Trader DNA still reuse the 12 demo exercises, so they show up as repeats there.", False),
    ("", False),
    ("1. Read the prompt, the question, the correct answer and both feedback texts.", False),
    ("2. Check the finance is right, the answer is unambiguous and nothing reads like financial advice or a promised return.", False),
    ("3. Set Status: OK, Change (write what in \"Reviewer comment\" or a new text in \"Suggested rewrite\") or Remove.", False),
    ("4. The Summary tab counts what's done per course.", False),
    ("", False),
    ("Text in [brackets] is filled in live from what the learner did (sizes, counts, prices).", False),
    ("Interactive exercises (charts, calculators) are easier to judge in the app: localhost:3430/courses/lab/<type>, e.g. /courses/lab/tap-engulfing.", False),
    (f"{len(rows)} questions across {len(courses)} courses · generated from branch feat/pilot-paywall.", False),
]
for text, bold in lines:
    hw.append([text])
    hw.cell(row=hw.max_row, column=1).font = Font(bold=bold, size=14 if bold else 11)
    hw.cell(row=hw.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
hw.column_dimensions["A"].width = 110
wb.active = 1
wb.save(OUT)
print(len(rows), "rows;", sum(len(r["also"]) for r in rows), "test repeats folded;", counters)

# rows for the review tool
if len(sys.argv) > 3:
    ids = [ws.cell(row=r, column=1).value for r in range(2, len(rows) + 2)]
    out = []
    for qid, row in zip(ids, rows):
        out.append({
            "id": qid, "course": row["course"], "courseId": row["courseId"], "unit": row["unit"], "lesson": row["lesson"],
            "lessonId": row["lessonId"], "isTest": row["isTest"], "step": row["n"], "kind": row["kind"], "type": LABELS.get(row["kind"], row["kind"]),
            "learn": row["mode"] == "learn", "eyebrow": row["eyebrow"], "prompt": row["prompt"], "question": row["q"], "answer": row["a"],
            "explanation": row["right"], "wrong": row["wrong"], "hint": row["hint"],
            "links": (row.get("help") or {}).get("links") or [], "also": row["also"],
        })
    json.dump(out, open(sys.argv[3], "w"), ensure_ascii=False, indent=0)
    print("rows json:", len(out))
